# -*- coding: utf-8 -*-
"""第三、四问创新扩展：公共层。

口径（本目录全部结果统一采用）：
  * 环境插值 = PCHIP（过全部 241 个观测点、不平滑、不外推），t>14400 s 保持末值；
  * 达标判据 = 全域节点最大干基含水率 < 0.15 kg/kg（未舍入），报告时刻取 60 s 网格 t60；
  * 几何 = 圆柱径向一维；Q4 收缩坐标 xi=r/R(t)，附件2 的 R(t) 用 PCHIP。

本模块只读 A题改进2/求解代码/solve.py，不修改它。
"""
from pathlib import Path
import sys, json, math
import numpy as np

ROOT = Path(__file__).resolve().parent.parent
RES = ROOT / 'results'
FIG = ROOT / 'figures'
RES.mkdir(exist_ok=True)
FIG.mkdir(exist_ok=True)

SOLVER_DIR = Path(__file__).resolve().parent / 'solver'          # 本地自带副本（自包含）
if not (SOLVER_DIR / 'solve.py').is_file():                       # 兜底：外部工作区
    SOLVER_DIR = Path(r'D:\数学建模国赛\A题\A题改进2\求解代码')
sys.path.insert(0, str(SOLVER_DIR))
import solve as S                      # noqa: E402
import hashlib                         # noqa: E402
import openpyxl                        # noqa: E402
SOLVER_SHA256 = hashlib.sha256((SOLVER_DIR / 'solve.py').read_bytes()).hexdigest()
from scipy.integrate import solve_ivp  # noqa: E402
from scipy.interpolate import PchipInterpolator  # noqa: E402

# ---------------- 物理/工艺常数（题面或附件给定） ----------------
R0 = 0.02                    # m
L_LEN = 0.25                 # m
H_T = 25.0                   # W/(m^2 K)
H_M = 8e-7                   # m/s
T_AMB_TAIL = 50.165          # degC
C_AMB_TAIL = 0.04986
X0 = 2.55                    # 初始干基含水率
RHO_W = 1000.0               # kg/m^3

RHO_D0 = (650.0 + 128.0 * X0) / (1.0 + X0)   # 附录3 初始干基体积密度 275.042 kg/m^3
M_DRY_PER_AREA = RHO_D0 * R0 / 2.0           # kg 干物质 / m^2 侧面积 = 2.7504

# ---------------- 品质降解动力学（文献量级，必须声明未标定） ----------------
R_GAS = 8.314
T_REF_K = 333.15             # 60 degC
K_REF_PER_H = 0.02           # 60 degC 参考降解速率常数 1/h
EA_REF = 50_000.0            # J/mol
EA_SET = (30_000.0, 50_000.0, 80_000.0)

# ---------------- 三段恒温制度 ----------------
SWITCH_H = (8.0, 24.0)
T_BOX = (40.0, 70.0)         # 优化区间（degC），超出附件观测范围，属工艺扩展假设
T_OBS_MAX = 50.246           # 附件1 观测温度最大值（用于"观测温区内"子问题）


# ==================== 模型：可设定三段恒温 ====================
class SchedModel(S.Model):
    """三段恒温制度；含湿量序列沿用附件1 观测（不随温度制度改变）。"""

    def __init__(self, case, n, temps, switch_h=SWITCH_H, use_obs=False, **kw):
        kw.setdefault('tail', 'last')
        kw.setdefault('mesh', 'surface')
        kw.setdefault('env_interp', 'pchip')
        super().__init__(case, n, **kw)
        self.temps = tuple(float(t) for t in temps)
        self.switch_s = tuple(float(h) * 3600.0 for h in switch_h)
        self.use_obs = bool(use_obs)            # True: 完全复现附件1 观测温度（机理性对照）

    def Ta(self, t):
        if self.use_obs:
            tc = float(np.clip(t, S.ENV[0, 0], S.ENV[-1, 0]))
            return float(self.env_curve(tc)[0])
        if t < self.switch_s[0]:
            return self.temps[0]
        if t < self.switch_s[1]:
            return self.temps[1]
        return self.temps[2]

    def env(self, t):
        tc = float(np.clip(t, S.ENV[0, 0], S.ENV[-1, 0]))
        ca = float(self.env_curve(tc)[1])       # 与基准完全相同的含湿量通道
        return self.Ta(t), ca


class CrustModel(SchedModel):
    """表皮阻力情景：表面传质系数按表面干基含水率衰减。"""

    def __init__(self, case, n, temps, k_crust=0.0, c_crust=0.5, **kw):
        super().__init__(case, n, temps, **kw)
        self.k_crust = float(k_crust)
        self.c_crust = float(c_crust)
        self.hm_base = float(self.hm)

    def rhs(self, t, y):
        m = self.n + 1
        cs = float(max(y[m:2 * m][-1], 0.0))
        if self.k_crust > 0.0:
            red = max(0.0, 1.0 - cs / self.c_crust)
            self.hm = self.hm_base / (1.0 + self.k_crust * red)
        else:
            self.hm = self.hm_base
        return super().rhs(t, y)


class TabRModel(S.Model):
    """Q4：使用外部给出的 R(t) 表（内生生收缩本构）代替附件2。"""

    def __init__(self, n, t_tab, R_tab, **kw):
        kw.setdefault('tail', 'last')
        kw.setdefault('mesh', 'surface')
        kw.setdefault('env_interp', 'pchip')
        super().__init__(4, n, **kw)
        self._fR = PchipInterpolator(t_tab, R_tab, extrapolate=True)
        self._tmax = float(t_tab[-1])

    def R(self, t):
        return float(max(self._fR(min(float(t), self._tmax)), 1e-4))


# ==================== 求解与后处理 ====================
def _extended(model, sol, t_end, rtol, max_step, method):
    """把解从事件点延长到 t_end，返回 (t数组, Y数组)。"""
    t = sol.t
    Y = sol.y
    if t_end > t[-1] + 1e-9:
        ext = solve_ivp(model.rhs, (t[-1], t_end), Y[:, -1], method=method,
                        rtol=rtol, atol=rtol * 0.02, max_step=max_step,
                        dense_output=True, jac=model.jacobian)
        if not ext.success:
            raise RuntimeError(ext.message)
        t = np.r_[t, ext.t[1:]]
        Y = np.c_[Y, ext.y[:, 1:]]
    return t, Y


def qoi(model, sol, ea=EA_REF, k_ref=K_REF_PER_H, rho_s=1400.0, want_traj=True,
        rtol=1e-6, max_step=300.0, method='BDF'):
    """从一次求解中提取全部关注的量。"""
    reached = bool(sol.t_events is not None and len(sol.t_events[0]) > 0)
    if not reached:
        return dict(reached=False, t_cross_s=None, t60_s=None, t60_h=None)
    tc = float(sol.t_events[0][0])
    t60 = 60.0 * math.floor(tc / 60.0 + 1.0)
    t, Y = _extended(model, sol, t60, rtol, max_step, method)
    mm = model.n + 1
    T = Y[:mm]                       # degC，节点
    C = Y[mm:2 * mm]                 # kg/kg，节点
    w = 2.0 * model.vol              # 体积权重，sum=1
    Xbar = w @ C
    TbarK = (w @ T) + 273.15
    Ts = T[-1]
    Ta = np.array([model.env(tt)[0] for tt in t])
    R = np.array([model.R(tt) for tt in t])

    q = H_T * (Ta - Ts)                                   # W/m^2 进入物料
    Q_heat = float(np.trapezoid(q, t))                    # J/m^2
    dX = float(X0 - Xbar[-1])
    water = dX * M_DRY_PER_AREA                           # kg/m^2
    sec = float(Q_heat / water) if water > 0 else float('inf')

    kdeg = (k_ref / 3600.0) * np.exp(-ea / R_GAS * (1.0 / TbarK - 1.0 / T_REF_K))   # 1/s
    dose = float(np.trapezoid(kdeg, t))
    qret = float(math.exp(-dose))

    rho_d = RHO_D0 * (R0 / R) ** 2
    phi_void = 1.0 - rho_d / rho_s
    phi_water = rho_d * Xbar / RHO_W
    phi_gas = phi_void - phi_water

    out = dict(reached=True, t_cross_s=tc, t60_s=t60, t60_h=t60 / 3600.0,
               Q_heat_J_m2=Q_heat, dX=dX, SEC_J_per_kg=sec,
               dose=dose, Q_ret=qret, Xbar_end=float(Xbar[-1]),
               Tbar_end_C=float(TbarK[-1] - 273.15), Ts_end_C=float(Ts[-1]),
               R_end_cm=float(R[-1] * 100), Cmax_end=float(C[:, -1].max()),
               rho_d_end=float(rho_d[-1]), phi_void_end=float(phi_void[-1]),
               phi_water_end=float(phi_water[-1]), phi_gas_end=float(phi_gas[-1]),
               min_phi_gas=float(np.min(phi_gas)))
    if want_traj:
        # 统一到 0.25 h 网格，便于跨样本比较与换 Ea 重算
        grid = np.arange(0.0, t60 + 1e-9, 900.0)
        grid = np.r_[grid[grid < t60], t60]
        out['traj'] = dict(
            t_h=(grid / 3600.0).tolist(),
            Xbar=np.interp(grid, t, Xbar).tolist(),
            TbarK=np.interp(grid, t, TbarK).tolist(),
            Ta=np.interp(grid, t, Ta).tolist(),
            Ts=np.interp(grid, t, Ts).tolist(),
            R_cm=(np.interp(grid, t, R) * 100).tolist(),
            q_W_m2=np.interp(grid, t, q).tolist(),
        )
    return out


def solve_sched(case, temps, n=160, rtol=1e-6, max_step=300.0, method='BDF',
                limit_h=150.0, switch_h=SWITCH_H, model_cls=SchedModel, **mk):
    model = model_cls(case, n, temps, switch_h=switch_h, **mk)
    sol = model.solve(limit_h * 3600.0, rtol=rtol, max_step=max_step, method=method)
    return model, sol


def run_schedule(case, temps, n=160, rtol=1e-6, max_step=300.0, method='BDF',
                 limit_h=150.0, switch_h=SWITCH_H, ea=EA_REF, want_traj=True, **mk):
    model, sol = solve_sched(case, temps, n, rtol, max_step, method, limit_h, switch_h, **mk)
    d = qoi(model, sol, ea=ea, rtol=rtol, max_step=max_step, method=method, want_traj=want_traj)
    d['temps'] = list(map(float, temps))
    d['switch_h'] = list(switch_h)
    d['n'] = n
    return d


def dose_from_traj(traj, ea, k_ref=K_REF_PER_H):
    """用已保存的 Tbar 轨迹重算任意 Ea 下的降解剂量（不需要重新求解 PDE）。"""
    t = np.asarray(traj['t_h']) * 3600.0
    Tb = np.asarray(traj['TbarK'])
    k = (k_ref / 3600.0) * np.exp(-ea / R_GAS * (1.0 / Tb - 1.0 / T_REF_K))
    return float(np.trapezoid(k, t))


def jdump(obj, name):
    p = RES / name
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=float), encoding='utf-8')
    return p


# ==================== 场采样与 result 工作簿 ====================
RADII_CM = np.round(np.arange(0.0, 2.0001, 0.1), 1)


def _t60_and_extend(model, sol, rtol, max_step, method):
    reached = bool(sol.t_events is not None and len(sol.t_events[0]) > 0)
    if not reached:
        raise RuntimeError('积分上限内未达标，不能导出达标结果')
    tc = float(sol.t_events[0][0])
    t60 = 60.0 * math.floor(tc / 60.0 + 1.0)
    t, Y = _extended(model, sol, t60, rtol, max_step, method)
    return tc, t60, t, Y


def sample_C(model, t_arr, Y, times, radii_cm=RADII_CM):
    """把节点解场采样到固定距离列；Q4 域外为 NaN。返回 (C矩阵, 表面列)。"""
    m = model.n + 1
    C = np.full((len(times), len(radii_cm)), np.nan)
    surf = np.full(len(times), np.nan)
    for i, tt in enumerate(times):
        j = int(np.searchsorted(t_arr, tt))
        j = min(max(j, 0), len(t_arr) - 1)
        x = model.x
        R = model.R(tt)
        r_or_xi = np.minimum(radii_cm / (R * 100.0), 1.0)
        vals = Y[m:2 * m, j]
        inside = radii_cm <= R * 100.0 + 2e-12
        C[i, inside] = np.interp(r_or_xi[inside], x, vals)
        surf[i] = vals[-1]
    return C, surf


def sample_field(model, sol, rtol=1e-7, max_step=200.0, method='BDF', dt=60.0):
    """返回 60 s 网格（首点=60 s，末点=t60）上的 21 距离列与表面列。"""
    tc, t60, t_arr, Y = _t60_and_extend(model, sol, rtol, max_step, method)
    times = np.arange(dt, t60 + 1e-9, dt)
    C, surf = sample_C(model, t_arr, Y, times)
    R = np.array([model.R(tt) for tt in times])
    return dict(t_cross_s=tc, t60_s=t60, times=times, C=C, surface=surf,
                R_cm=R * 100.0, t_arr=t_arr, Y=Y)


def write_single_sheet(path, times, C, surface=None, ndigits=4, sheet='Sheet1'):
    """C: (k,21)。surface 非空时追加"药材表面"列（result4 式）。"""
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(sheet)
    hdr = ['时间\\到药材中心的距离'] + [f'{c:g}' for c in RADII_CM] + \
          (['药材表面'] if surface is not None else [])
    ws.append(hdr)
    for i, t in enumerate(times):
        row = [int(round(float(t)))]
        row += [None if not np.isfinite(v) else round(float(v), ndigits) for v in C[i]]
        if surface is not None:
            v = surface[i]
            row.append(None if not np.isfinite(v) else round(float(v), ndigits))
        ws.append(row)
    wb.save(path)


def write_two_sheet(path, times, C, T, ndigits=4):
    """温度/水分浓度两张工作表（result1/result2 式）。"""
    wb = openpyxl.Workbook(write_only=True)
    hdr = ['时间\\到药材中心的距离'] + [f'{c:g}' for c in RADII_CM]
    for name, mat in (('温度', T), ('水分浓度', C)):
        ws = wb.create_sheet(name)
        ws.append(hdr)
        for i, t in enumerate(times):
            ws.append([int(round(float(t)))] +
                      [None if not np.isfinite(v) else round(float(v), ndigits) for v in mat[i]])
    wb.save(path)


def sample_T(model, t_arr, Y, times):
    """同 sample_C，但取温度场（℃）。"""
    m = model.n + 1
    T = np.full((len(times), len(RADII_CM)), np.nan)
    for i, tt in enumerate(times):
        j = min(max(int(np.searchsorted(t_arr, tt)), 0), len(t_arr) - 1)
        R = model.R(tt)
        r_or_xi = np.minimum(RADII_CM / (R * 100.0), 1.0)
        vals = Y[:m, j]
        inside = RADII_CM <= R * 100.0 + 2e-12
        T[i, inside] = np.interp(r_or_xi[inside], model.x, vals)
    return T


def build_sched_model(case, temps, n=320, switch_h=SWITCH_H, use_obs=False, k_crust=0.0, **kw):
    kw.setdefault('tail', 'last')
    kw.setdefault('mesh', 'surface')
    kw.setdefault('env_interp', 'pchip')
    if k_crust > 0:
        return CrustModel(case, n, temps, k_crust=k_crust, use_obs=use_obs, switch_h=switch_h, **kw)
    return SchedModel(case, n, temps, switch_h=switch_h, use_obs=use_obs, **kw)


def build_tabR_model(t_tab, R_tab, n=320, **kw):
    kw.setdefault('tail', 'last')
    kw.setdefault('mesh', 'surface')
    kw.setdefault('env_interp', 'pchip')
    return TabRModel(n, t_tab, R_tab, **kw)


def q3_baseline_model(n=320):
    """题目口径基准：附件1 观测环境（PCHIP）。"""
    return S.Model(3, n, tail='last', mesh='surface', env_interp='pchip')


# ==================== 基准口径 ====================
def baseline():
    """题目口径基准：附件1 观测环境（PCHIP），Q3/Q4 各一次。"""
    out = {}
    for case in (3, 4):
        model = S.Model(case, 160, tail='last', mesh='surface', env_interp='pchip')
        sol = model.solve(150 * 3600.0, rtol=1e-6, max_step=300.0, method='BDF')
        out[f'q{case}'] = qoi(model, sol, want_traj=True, rtol=1e-6, max_step=300.0)
    return out


if __name__ == '__main__':
    print(json.dumps({k: v for k, v in baseline().items()}, ensure_ascii=False, indent=2,
                     default=float))
