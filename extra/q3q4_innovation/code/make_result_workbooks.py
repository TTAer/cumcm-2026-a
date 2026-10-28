# -*- coding: utf-8 -*-
"""把第三、四问的创新结果落成**官方 result 格式**的工作簿。

产出（results/result_创新/）：
  Q3_最速_result3式.xlsx          Q3 时间最优三段制度下的 60 s 网格含水率场
  Q3_TOPSIS折中_result3式.xlsx     Q3 三目标折中制度
  Q3_最省能_result3式.xlsx         Q3 低温长时制度
  Q3_观测上限_result3式.xlsx       温度不超过附件1 观测上限的退化情形
  Q4_附件2口径_result4式.xlsx      Q4 题面口径（附件2 半径）
  Q4_内生二次_result4式.xlsx       Q4 把附件2 换成内生二次收缩本构
  Q4_表皮Kc2_result4式.xlsx        Q4 表皮阻力情景
  创新工艺汇总.xlsx                Pareto 前沿与推荐制度总表（非模板，另存）
另附（主答案 P0 修复，B-02）：
  result2_前3h.xlsx                Q2 主表：1–10800 s 每秒、温度+水分浓度双表
  result2_全程60s辅助表.xlsx        同一连续积分的 60 s 全程辅助表

**注意：本目录全部工作簿都是"扩展讨论"产物，不替代题目主答案
（Q3 = 205860 s、Q4 = 183000 s），文件名与存放目录均与正式 result 区分。**
"""
import json, sys, math
from pathlib import Path
import numpy as np
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common as K   # noqa: E402

OUT = K.RES / 'result_创新'
OUT.mkdir(parents=True, exist_ok=True)
N = 320
RTOL = 1e-7
MAX_STEP = 200.0


def solve_field(model, limit_h=150.0):
    sol = model.solve(limit_h * 3600.0, rtol=RTOL, max_step=MAX_STEP, method='BDF')
    f = K.sample_field(model, sol, rtol=RTOL, max_step=MAX_STEP)
    f['model'] = model
    return f


FIELDS = {}          # 供画图复用：name -> (times, C, surface)


def q3_case(path, temps, label, use_obs=False, n=N, extra=None):
    m = K.build_sched_model(3, temps, n=n, use_obs=use_obs)
    f = solve_field(m)
    K.write_single_sheet(OUT / path, f['times'], f['C'])
    FIELDS[path] = dict(times=f['times'], C=f['C'], kind='q3', temps=list(temps), label=label)
    rec = dict(label=label, file=path, temps=list(temps), switch_h=list(K.SWITCH_H),
               n=n, rtol=RTOL, max_step=MAX_STEP, use_obs=use_obs,
               t_cross_s=f['t_cross_s'], t60_s=f['t60_s'], t60_h=f['t60_s'] / 3600.0,
               rows=int(len(f['times'])), first_time=float(f['times'][0]),
               last_time=float(f['times'][-1]),
               C_end_max=float(np.nanmax(f['C'][-1])),
               g_prev=float(np.nanmax(f['C'][-2]) - 0.15),
               g_stop=float(np.nanmax(f['C'][-1]) - 0.15),
               bytes=(OUT / path).stat().st_size)
    if extra:
        rec.update(extra)
    print(f'  {label:<16} {path:<34} t60={rec["t60_h"]:.4f} h  行数={rec["rows"]}  '
          f'{rec["bytes"]/1024:.0f} KB', flush=True)
    return rec


def q4_case(path, model, label, r_src):
    f = solve_field(model)
    K.write_single_sheet(OUT / path, f['times'], f['C'], surface=f['surface'])
    FIELDS[path] = dict(times=f['times'], C=f['C'], surface=f['surface'], R_cm=f['R_cm'],
                        kind='q4', label=label)
    nblank = int(np.isnan(f['C']).sum())
    rec = dict(label=label, file=path, R_source=r_src, n=N, rtol=RTOL, max_step=MAX_STEP,
               t_cross_s=f['t_cross_s'], t60_s=f['t60_s'], t60_h=f['t60_s'] / 3600.0,
               rows=int(len(f['times'])), first_time=float(f['times'][0]),
               last_time=float(f['times'][-1]),
               blanks=nblank, R_end_cm=float(f['R_cm'][-1]),
               surface_end=float(f['surface'][-1]),
               C_end_max=float(np.nanmax(f['C'][-1])),
               g_prev=float(np.nanmax(f['C'][-2]) - 0.15),
               g_stop=float(np.nanmax(f['C'][-1]) - 0.15),
               bytes=(OUT / path).stat().st_size)
    print(f'  {label:<16} {path:<34} t60={rec["t60_h"]:.4f} h  行数={rec["rows"]}  '
          f'域外空格={nblank}  {rec["bytes"]/1024:.0f} KB', flush=True)
    return rec


def main():
    q3 = json.loads((K.RES / 'q3_优化结果.json').read_text(encoding='utf-8'))
    q4 = json.loads((K.RES / 'q4_收缩与孔隙率.json').read_text(encoding='utf-8'))
    recs = []

    # ---------- Q3 四个制度 ----------
    print('Q3：按官方 result3 格式导出三个推荐制度与观测上限制度', flush=True)
    picks = []
    seen = set()
    for key, lab in (('min_time', '最速'), ('topsis', 'TOPSIS折中'), ('min_energy', '最省能')):
        r = q3['recommend'][key]
        tp = tuple(round(x, 2) for x in r['temps'])
        if tp in seen:
            continue
        seen.add(tp)
        picks.append((f'Q3_{lab}_result3式.xlsx', r['temps'], lab, r['Q_ret']))
    obs = [r for r in q3['pareto_verified'] if r['src'] == 'obs_upper'][0]
    picks.append(('Q3_观测上限_result3式.xlsx', obs['temps'], '观测上限恒温', obs['Q_ret']))

    for fn, temps, lab, qret in picks:
        recs.append(q3_case(fn, temps, lab, extra=dict(Q_ret=float(qret))))

    # ---------- Q4 三个口径 ----------
    print('Q4：按官方 result4 格式导出三个收缩/表皮口径', flush=True)
    mb = K.S.Model(4, N, tail='last', mesh='surface', env_interp='pchip')
    fb = solve_field(mb)
    recs.append(q4_case('Q4_附件2口径_result4式.xlsx', mb,
                        '附件2 半径（题面口径）', '附件2 PCHIP'))

    # 内生二次：用基准 Xbar 轨迹经二次本构重建 R(t)
    mm4 = mb.n + 1
    Xb = (2.0 * mb.vol) @ fb['Y'][mm4:2 * mm4]
    dX = Xb - K.X0
    par = q4['constitutive']['forms']['径向二次_过初值']['params']
    rr = 1.0 + par['k1'] * dX + par['k2'] * dX ** 2
    Rb = K.R0 * np.minimum.accumulate(np.clip(rr, 0.2, 1.0))
    mt = K.build_tabR_model(fb['t_arr'], Rb, n=N)
    recs.append(q4_case('Q4_内生二次_result4式.xlsx', mt, '内生二次收缩本构',
                        f"R/R0 = 1 + {par['k1']:.6f}(X-X0) + {par['k2']:.6f}(X-X0)^2"))

    # 表皮阻力 K_c = 2
    mc = K.build_sched_model(4, (50.165,) * 3, n=N, use_obs=True, k_crust=2.0)
    recs.append(q4_case('Q4_表皮Kc2_result4式.xlsx', mc, '表皮阻力 K_c=2（C_crust=0.5）',
                        '附件2 PCHIP + 表面传质衰减'))

    # ---------- 汇总工作簿 ----------
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '推荐制度'
    ws.append(['类别', '制度', 'T1/℃', 'T2/℃', 'T3/℃', 't60/h', 't60/s',
               'SEC/(MJ/kg)', 'SEC+λ/(MJ/kg)', 'Q_ret', '峰值热流/(W/m²)', '相对题面基准/h'])
    base_t = q3['baseline']['q3_t60_h']
    for key, lab in (('min_time', '最速（时间最优）'), ('topsis', 'TOPSIS 折中'),
                     ('min_energy', '最省能'), ('max_quality', '最优品质')):
        r = q3['recommend'][key]
        if not isinstance(r, dict) or 'temps' not in r:
            continue
        ws.append(['Q3 三段恒温', lab, round(r['temps'][0], 2), round(r['temps'][1], 2),
                   round(r['temps'][2], 2), round(r['t60_h'], 4), int(round(r['t60_h'] * 3600)),
                   round(r['SEC_MJ'], 4), round(r['SEC_MJ'] + q3['lam_vap_J_per_kg'] / 1e6, 4),
                   round(r['Q_ret'], 4), round(r['qmax_W_m2'], 1), round(r['t60_h'] - base_t, 4)])
    ws.append(['Q3 基准', '附件1 观测环境（题面口径）', '—', '—', '—', round(base_t, 4),
               int(round(base_t * 3600)), round(q3['baseline']['q3_SEC_MJ'], 4),
               round(q3['baseline']['q3_SEC_MJ'] + q3['lam_vap_J_per_kg'] / 1e6, 4),
               round(q3['baseline']['q3_Qret'], 4), '—', 0.0])
    ws.append([])
    ws.append(['类别', '收缩/传质口径', '', '', '', 't60/h', 't60/s', '', '', '', '', '相对附件2 口径/h'])
    for kk, v in q4['endogenous'].items():
        ws.append(['Q4 收缩口径', kk.replace('_', ' '), '', '', '', round(v['t60_h'], 4),
                   int(round(v['t60_h'] * 3600)), '', '', '', '', round(v['delta_h'], 4)])
    for kk, v in q4['crust'].items():
        ws.append(['Q4 表皮阻力', f'K_c={kk}', '', '', '', round(v['t60_h'], 4),
                   int(round(v['t60_h'] * 3600)), '', '', '', '', round(v['delta_h'], 4)])
    ws2 = wb.create_sheet('Pareto前沿')
    ws2.append(['T1/℃', 'T2/℃', 'T3/℃', 't60/h', 'SEC/(MJ/kg)', 'Q_ret',
                '峰值热流/(W/m²)', '物料最高温/℃', '来源'])
    full = {tuple(round(x, 6) for x in r['temps']): r for r in q3['pareto_verified']}
    for r in sorted(q3['pareto_verified_front'], key=lambda x: x['t60_h']):
        fr = full.get(tuple(round(x, 6) for x in r['temps']), {})
        ws2.append([round(r['temps'][0], 2), round(r['temps'][1], 2), round(r['temps'][2], 2),
                    round(r['t60_h'], 4), round(r['SEC_MJ'], 4), round(r['Q_ret'], 4),
                    round(fr.get('qmax_W_m2', float('nan')), 1),
                    round(fr.get('Tbar_max_C', float('nan')), 2), r['src']])
    ws3 = wb.create_sheet('说明')
    for line in [
        '本工作簿收录第三、四问创新扩展的计算结果，不替代题目主答案。',
        f'题目主答案（同一判据口径）：Q3 t60 = 205860 s = {base_t:.4f} h；Q4 t60 = 183000 s = 50.8333 h。',
        'Q3 三段恒温切换时点固定 8 h / 24 h；温度区间 40–70 ℃ 超出附件1 观测范围（28–50.246 ℃），属工艺扩展假设。',
        'SEC 只统计对流供入物料的热量（本模型不含潜热项）；SEC+λ 为可比口径（λ = 2.383 MJ/kg）。',
        '品质动力学 k_ref = 0.02 h⁻¹ @60 ℃、Ea = 50 kJ/mol 取自文献量级，未标定，Q_ret 仅作相对排序。',
        'Q4 内生生收缩本构使用 A题改进4 的径向二次拟合；附件2 与附录4 密度式不相容，详见 results/关键结果汇总.md。',
        '全部数据由 code/run_all.py 与 code/make_result_workbooks.py 一键复现。',
    ]:
        ws3.append([line])
    wb.save(OUT / '创新工艺汇总.xlsx')
    print(f'  汇总工作簿 创新工艺汇总.xlsx  {(OUT / "创新工艺汇总.xlsx").stat().st_size/1024:.0f} KB',
          flush=True)

    # ---------- 主答案 P0 修复：result2 时域 ----------
    print('主答案 P0 修复（B-02）：result2 前 3 h 主表 + 60 s 全程辅助表', flush=True)
    m2 = K.q3_baseline_model(n=N)          # Q2 与 Q3 共享同一次连续积分
    sol2 = m2.solve(150 * 3600.0, rtol=RTOL, max_step=MAX_STEP, method='BDF')
    tc, t60, t_arr, Y = K._t60_and_extend(m2, sol2, RTOL, MAX_STEP, 'BDF')
    t1 = np.arange(1.0, 10800.0 + 1e-9, 1.0)
    C1 = K.sample_C(m2, t_arr, Y, t1)[0]
    T1 = K.sample_T(m2, t_arr, Y, t1)
    K.write_two_sheet(OUT / 'result2_前3h.xlsx', t1, C1, T1)
    t60g = np.arange(60.0, t60 + 1e-9, 60.0)
    C2 = K.sample_C(m2, t_arr, Y, t60g)[0]
    T2 = K.sample_T(m2, t_arr, Y, t60g)
    K.write_two_sheet(OUT / 'result2_全程60s辅助表.xlsx', t60g, C2, T2)
    FIELDS['result2_前3h.xlsx'] = dict(times=t1, C=C1, T=T1, kind='q2', label='题面观测环境（Q2/Q3 共享积分）')
    FIELDS['result2_全程60s辅助表.xlsx'] = dict(times=t60g, C=C2, T=T2, kind='q23full',
                                                label='题面观测环境（全程 60 s）')
    r2 = dict(前3h=dict(file='result2_前3h.xlsx', rows=int(len(t1)), first=1, last=10800,
                        bytes=(OUT / 'result2_前3h.xlsx').stat().st_size),
              全程60s=dict(file='result2_全程60s辅助表.xlsx', rows=int(len(t60g)), first=60,
                           last=int(t60), bytes=(OUT / 'result2_全程60s辅助表.xlsx').stat().st_size),
              t60_s=t60, t_cross_s=tc)
    print(f'  result2_前3h.xlsx 行数={len(t1)}  {r2["前3h"]["bytes"]/1024/1024:.2f} MB', flush=True)
    print(f'  result2_全程60s辅助表.xlsx 行数={len(t60g)}  {r2["全程60s"]["bytes"]/1024/1024:.2f} MB',
          flush=True)

    man = dict(
        note='第三、四问创新扩展的官方格式结果工作簿；不替代题目主答案。',
        solver_sha256=K.SOLVER_SHA256,
        config=dict(n=N, rtol=RTOL, max_step=MAX_STEP, radii_cm=K.RADII_CM.tolist(),
                    switch_h=list(K.SWITCH_H), box=list(K.T_BOX)),
        main_answer=dict(q3_t60_s=205860, q3_t60_h=205860 / 3600,
                         q4_t60_s=183000, q4_t60_h=183000 / 3600),
        q3_cases=[r for r in recs if r['file'].startswith('Q3')],
        q4_cases=[r for r in recs if r['file'].startswith('Q4')],
        result2_fix=r2,
        total_bytes=sum(p.stat().st_size for p in OUT.glob('*.xlsx')),
    )
    K.jdump(man, 'result_创新/清单.json')

    # 场数据另存 NPZ，供画图与二次分析复用（不重复求解）
    np.savez_compressed(
        K.RES / '创新场.npz',
        names=np.array(list(FIELDS.keys())),
        **{f'{k}__t': v['times'] for k, v in FIELDS.items()},
        **{f'{k}__C': v['C'] for k, v in FIELDS.items()},
        **{f'{k}__S': v['surface'] for k, v in FIELDS.items() if 'surface' in v},
        **{f'{k}__R': v['R_cm'] for k, v in FIELDS.items() if 'R_cm' in v},
        **{f'{k}__T': v['T'] for k, v in FIELDS.items() if 'T' in v})
    print(f'  场数据 创新场.npz  {(K.RES / "创新场.npz").stat().st_size/1024/1024:.2f} MB', flush=True)

    print(f'\n已写出 {OUT}  共 {len(list(OUT.glob("*.xlsx")))} 个工作簿，'
          f'{man["total_bytes"]/1024/1024:.2f} MB', flush=True)


if __name__ == '__main__':
    main()
